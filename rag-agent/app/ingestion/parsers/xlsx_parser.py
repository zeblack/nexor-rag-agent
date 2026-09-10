import os

from app.ingestion.parsers import ParsedDocument


def parse(file_path: str) -> ParsedDocument:
    ext = os.path.splitext(file_path)[1].lower()
    if ext == ".xlsb":
        return _parse_xlsb(file_path)
    return _parse_xlsx(file_path)


def _parse_xlsx(file_path: str) -> ParsedDocument:
    import openpyxl

    wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    parts = []
    try:
        for sheet in wb.worksheets:
            parts.append(f"## {sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                if any(c is not None for c in row):
                    parts.append(" | ".join("" if c is None else str(c) for c in row))
    finally:
        wb.close()
    return ParsedDocument(text="\n".join(parts), parser_used="openpyxl")


def _parse_xlsb(file_path: str) -> ParsedDocument:
    import pyxlsb

    parts = []
    with pyxlsb.open_workbook(file_path) as wb:
        for sheet_name in wb.sheets:
            parts.append(f"## {sheet_name}")
            with wb.get_sheet(sheet_name) as sheet:
                for row in sheet.rows():
                    values = [item.v for item in row]
                    if any(v is not None for v in values):
                        parts.append(" | ".join("" if v is None else str(v) for v in values))
    return ParsedDocument(text="\n".join(parts), parser_used="pyxlsb")
