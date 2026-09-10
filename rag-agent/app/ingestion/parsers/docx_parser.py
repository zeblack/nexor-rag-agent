import docx

from app.ingestion.parsers import ParsedDocument


def parse(file_path: str) -> ParsedDocument:
    document = docx.Document(file_path)
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return ParsedDocument(text="\n".join(parts), parser_used="python-docx")
