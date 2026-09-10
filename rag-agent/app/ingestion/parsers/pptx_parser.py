from pptx import Presentation

from app.ingestion.parsers import ParsedDocument


def parse(file_path: str) -> ParsedDocument:
    prs = Presentation(file_path)
    parts = []
    for i, slide in enumerate(prs.slides, start=1):
        parts.append(f"## Slide {i}")
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                parts.append(shape.text_frame.text)
    return ParsedDocument(text="\n".join(parts), parser_used="python-pptx")
